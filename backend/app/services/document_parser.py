import re
from dataclasses import dataclass
from pathlib import Path

import fitz
from docx import Document

from app.core.config import settings
from app.core.constants import FileType


@dataclass
class ParsedDocument:
    text: str
    page_count: int | None
    word_count: int
    warnings: list[str]


class DocumentParserService:
    async def parse(self, path: Path, file_type: str) -> ParsedDocument:
        if file_type == FileType.PDF:
            return await self._parse_pdf(path)
        if file_type == FileType.DOCX:
            return self._parse_docx(path)
        if file_type == FileType.TXT:
            text = path.read_text(encoding="utf-8", errors="ignore")
            return ParsedDocument(self._normalize(text), None, len(text.split()), [])
        raise ValueError(f"Unsupported file type: {file_type}")

    async def _parse_pdf(self, path: Path) -> ParsedDocument:
        warnings: list[str] = []
        doc = fitz.open(path)
        page_texts = [page.get_text("text") for page in doc]
        text = self._normalize("\n\n".join(page_texts))
        if len(text.split()) < max(30, len(doc) * 5):
            warnings.append("PDF_TEXT_SPARSE")
            ocr_text = self._ocr_pdf(path) if settings.ocr_enabled else ""
            if ocr_text:
                text = self._normalize(ocr_text)
                warnings.append("OCR_FALLBACK_USED")
            else:
                warnings.append("OCR_FALLBACK_NOT_AVAILABLE")
        return ParsedDocument(text=text, page_count=len(doc), word_count=len(text.split()), warnings=warnings)

    def _parse_docx(self, path: Path) -> ParsedDocument:
        document = Document(path)
        blocks: list[str] = []
        blocks.extend(p.text for p in document.paragraphs if p.text.strip())
        for table in document.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
        text = self._normalize("\n".join(blocks))
        return ParsedDocument(text=text, page_count=None, word_count=len(text.split()), warnings=[])

    def _ocr_pdf(self, path: Path) -> str:
        try:
            from pdf2image import convert_from_path
            import pytesseract
        except Exception:
            return ""
        images = convert_from_path(str(path), dpi=200)
        return "\n\n".join(pytesseract.image_to_string(image) for image in images)

    def _normalize(self, text: str) -> str:
        text = text.replace("\x00", " ")
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()
